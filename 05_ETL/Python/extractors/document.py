from pathlib import Path

import docx

from utils import normalize_spaces


def get_table_as_dict(doc, table_index=0):
    """
    Convertit un tableau Word de type clé / valeur en dictionnaire.

    La première cellule de chaque ligne devient la clé,
    la deuxième cellule devient la valeur.
    """
    result = {}

    if len(doc.tables) <= table_index:
        return result

    table = doc.tables[table_index]

    for row in table.rows:
        cells = [
            normalize_spaces(cell.text)
            for cell in row.cells
        ]

        if len(cells) < 2:
            continue

        key = cells[0]
        value = cells[1]

        if key:
            result[key] = value

    return result


def read_document(path: Path):
    """
    Lit un document Word et prépare les différentes représentations
    utilisées par les extracteurs métier.

    Retourne :
        - doc       : objet python-docx
        - paragraphs: paragraphes nettoyés
        - full_text : texte complet paragraphes + tableaux
        - meta      : métadonnées du premier tableau
    """

    doc = docx.Document(path)

    paragraphs = [
        normalize_spaces(p.text)
        for p in doc.paragraphs
        if p.text.strip()
    ]

    table_text = []

    for table in doc.tables:
        for row in table.rows:
            table_text.append(
                " | ".join(
                    normalize_spaces(cell.text)
                    for cell in row.cells
                )
            )

    full_text = "\n".join(paragraphs + table_text)

    meta = get_table_as_dict(doc, 0)

    return {
        "doc": doc,
        "paragraphs": paragraphs,
        "full_text": full_text,
        "meta": meta,
    }